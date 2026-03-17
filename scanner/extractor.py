import os
from pypdf import PdfReader
from docx import Document
from openpyxl import load_workbook
from PIL import Image
import exifread


def extract_metadata(file_path):
    ext = os.path.splitext(file_path)[1].lower()

    if ext == ".pdf":
        return extract_pdf_metadata(file_path)

    elif ext == ".docx":
        return extract_docx_metadata(file_path)

    elif ext == ".xlsx":
        return extract_xlsx_metadata(file_path)

    elif ext in [".jpg", ".jpeg", ".png"]:
        return extract_image_metadata(file_path)

    else:
        return {"error": "Unsupported file type"}


# ---------------- PDF ----------------
def extract_pdf_metadata(path):
    reader = PdfReader(path)
    meta = reader.metadata or {}

    return normalize_metadata({
        "author": meta.get("/Author"),
        "creator": meta.get("/Creator"),
        "producer": meta.get("/Producer"),
        "created": meta.get("/CreationDate"),
        "modified": meta.get("/ModDate"),
    })


# ---------------- DOCX ----------------
def extract_docx_metadata(path):
    doc = Document(path)
    core = doc.core_properties

    return normalize_metadata({
        "author": core.author,
        "last_modified_by": core.last_modified_by,
        "created": core.created,
        "modified": core.modified,
        "title": core.title,
    })


# ---------------- XLSX ----------------
def extract_xlsx_metadata(path):
    wb = load_workbook(path)
    props = wb.properties

    return normalize_metadata({
        "creator": props.creator,
        "last_modified_by": props.lastModifiedBy,
        "created": props.created,
        "modified": props.modified,
        "title": props.title,
    })


# ---------------- IMAGE ----------------
def extract_image_metadata(path):
    data = {}

    with open(path, 'rb') as img:
        tags = exifread.process_file(img, details=False)

        for tag in tags:
            data[tag] = str(tags[tag])

    return normalize_metadata(data)


# ---------------- NORMALIZER ----------------
def normalize_metadata(raw):
    cleaned = {}

    for k, v in raw.items():
        if v:
            cleaned[k.replace("/", "").lower()] = str(v)

    return cleaned
